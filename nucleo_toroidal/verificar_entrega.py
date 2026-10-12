"""Inspeciona documentos, malhas e links e empacota uma entrega local verificável."""
from pathlib import Path
from html.parser import HTMLParser
from xml.etree import ElementTree as ET
import hashlib
import json
import struct
import zipfile
import pymupdf
from docx import Document

ROOT=Path(__file__).resolve().parent

class Links(HTMLParser):
    def __init__(self): super().__init__(); self.links=[]
    def handle_starttag(self,tag,attrs):
        for name,val in attrs:
            if name in ('href','src') and val and not val.startswith(('http:','https:','#','data:')):self.links.append(val)

def main():
    pdf=pymupdf.open(ROOT/'Projeto_Tecnico.pdf')
    text='\n'.join(p.get_text() for p in pdf)
    needed=['174,392','78,405','1.041,400','44,3','84,6','Não existe','referências']
    required_tokens=['174,392','78,405','1.041,400','44,3','84,6','Ainda']
    numbers_present=all(x in text for x in required_tokens[:5])
    assert numbers_present,'Números de referência ausentes no PDF'
    assert len(pdf)>=5 and len(text)>12000,'PDF incompleto'
    pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.2,1.2)).save(ROOT/'Previa_Projeto_PDF.png')
    # Text extraction is also useful for independent review; it is not evidence of a physical test.
    (ROOT/'Projeto_Tecnico.txt').write_text(text,encoding='utf-8')
    parser=Links();parser.feed((ROOT/'index.html').read_text(encoding='utf-8'))
    missing=[link for link in parser.links if not (ROOT/link).exists()]
    assert not missing,missing
    mesh=json.loads((ROOT/'cad/validacao_mesh.json').read_text(encoding='utf-8'))
    assert all(m['boundary_edges']==0 and m['nonmanifold_edges']==0 and m['consistent_orientation'] and m['volume_mm3']>0 for m in mesh.values())
    stls=[]
    for f in sorted((ROOT/'cad').glob('*.stl')):
        b=f.read_bytes();n=struct.unpack_from('<I',b,80)[0];assert len(b)==84+50*n
        stls.append({'file':str(f.relative_to(ROOT)),'triangle_count':n,'binary_size_valid':True})
    for svg in [ROOT/'cad/desenho_dimensional.svg',ROOT/'circuito_ensaio.svg']:ET.parse(svg)
    document=Document(ROOT/'Projeto_Tecnico.docx')
    assert len(document.tables)>=4 and len(document.paragraphs)>60
    physics=json.loads((ROOT/'simulacao/resultados/resultados_validacao.json').read_text(encoding='utf-8'))
    browser=json.loads((ROOT/'validacao_visualizador.json').read_text(encoding='utf-8'))
    assert physics['all_passed'] and browser['all_passed']
    checks={
        'date':'2026-10-12','delivery_scope':'Proposta de bancada eletromagnética clássica; sem fabricação ou ensaio físico.',
        'physics_checks_passed':physics['passed_count'],'browser_checks_passed':browser['check_count'],
        'closed_mesh_geometry_types':len(mesh),'pdf_pages':len(pdf),'pdf_text_characters':len(text),
        'pdf_baseline_numbers_present':numbers_present,'docx_tables':len(document.tables),'missing_viewer_links':missing,
        'stl_files':stls,'svg_xml_valid':True,
        'not_performed':['Fabricação e montagem','Ensaio de campo/potência/temperatura','Renderização CLI do OpenSCAD','Certificação de equipamento','Comprovação quântica/gravitacional'],
        'all_digital_checks_passed':True,
    }
    (ROOT/'validacao_entrega.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    files=[p for p in sorted(ROOT.rglob('*')) if p.is_file() and '__pycache__' not in p.parts]
    sums=['# SHA-256 dos arquivos da entrega. validacao_entrega.json registra escopo/limites.']
    for file in files:
        if file.name=='SHA256SUMS.txt':continue
        sums.append(hashlib.sha256(file.read_bytes()).hexdigest()+'  '+file.relative_to(ROOT).as_posix())
    (ROOT/'SHA256SUMS.txt').write_text('\n'.join(sums)+'\n',encoding='utf-8')
    archive=ROOT.parent/'Nucleo_Toroidal_Pacote.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=8) as z:
        for p in sorted(ROOT.rglob('*')):
            if p.is_file() and '__pycache__' not in p.parts:z.write(p,arcname=Path('nucleo_toroidal')/p.relative_to(ROOT))
    with zipfile.ZipFile(archive) as z:assert z.testzip() is None
    print(json.dumps({'digital_checks':'passed','pdf_pages':len(pdf),'physics_checks':physics['passed_count'],'browser_checks':browser['check_count'],'mesh_types':len(mesh),'archive_bytes':archive.stat().st_size}))

if __name__=='__main__':main()
