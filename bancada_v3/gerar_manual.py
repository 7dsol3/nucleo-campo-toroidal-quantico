"""Consolidate V3 source documents into readable PDF and DOCX; no hardware claims.

Run from repository root: py -X utf8 -B bancada_v3/gerar_manual.py
Dependencies: reportlab, python-docx, PyMuPDF. Source Markdown remains canonical.
Only writes Manual_Bancada_V3.pdf and Manual_Bancada_V3.docx next to this script.
"""
from __future__ import annotations

import csv
from datetime import datetime, timezone
import hashlib
import html
import json
from pathlib import Path
import re
import textwrap
from urllib.parse import quote, urljoin
import zipfile

import pymupdf as fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Image, LongTable, PageBreak, Paragraph,
                               SimpleDocTemplate, Spacer, TableStyle)

BASE = Path(__file__).resolve().parent
REPO = BASE.parent
GITHUB = 'https://github.com/7dsol3/nucleo-campo-toroidal-quantico/blob/main/'
SOURCES = [
    ('README.md', '1. Visão geral, entregas e modos de operação'),
    ('eletronica/README.md', '2. Eletrônica, equipamentos e envelope calculado'),
    ('eletronica/LIGACOES.md', '3. Ligações, pinagem e fronteiras de potência'),
    ('eletronica/QUALIFICACAO.md', '4. Qualificação elétrica e térmica pendente'),
    ('eletronica/SENSOR_CAMPO.md', '5. Sensor de campo, temperatura e calibração'),
    ('aquisicao/README.md', '6. Captura simultânea PicoScope'),
    ('aquisicao/FILTRO_ANTIALIAS.md', '7. Filtro analógico e carregamento de medição'),
    ('medicoes/README.md', '8. Análise, proveniência e limites de erro'),
    ('controle/README.md', '9. Firmware PI CC de uma bobina'),
    ('fem/README.md', '10. FEM 3D executado e precisão documentada'),
    ('eletronica/FONTES.md', '11. Fontes primárias e parâmetros usados'),
]
LINK_PATTERN = re.compile(r'(!?)\[([^\]]+)\]\(([^)]+)\)')
DOCX_BODY_FONT = 'Arial' if Path('C:/Windows/Fonts/arial.ttf').exists() else 'DejaVu Sans'
DOCX_CODE_FONT = 'Consolas' if Path('C:/Windows/Fonts/consola.ttf').exists() else 'DejaVu Sans Mono'


class ManualDocTemplate(SimpleDocTemplate):
    def afterFlowable(self, flowable):
        title = getattr(flowable, 'manual_bookmark', None)
        if title:
            key = 'chapter_'+hashlib.sha256(title.encode('utf-8')).hexdigest()[:16]
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(title, key, level=0)


def font_setup():
    windows = Path('C:/Windows/Fonts')
    linux = Path('/usr/share/fonts/truetype/dejavu')
    names = [('Body', 'arial.ttf', 'DejaVuSans.ttf'),
             ('Body-Bold', 'arialbd.ttf', 'DejaVuSans-Bold.ttf'),
             ('Body-Italic', 'ariali.ttf', 'DejaVuSans-Oblique.ttf'),
             ('Body-BoldItalic', 'arialbi.ttf', 'DejaVuSans-BoldOblique.ttf'),
             ('Code', 'consola.ttf', 'DejaVuSansMono.ttf')]
    resolved = []
    for family, win, lin in names:
        path = windows/win if (windows/win).exists() else linux/lin
        if not path.exists():
            raise RuntimeError('TrueType font missing: Arial/Consolas or DejaVu Sans/Mono required')
        pdfmetrics.registerFont(TTFont(family, str(path)))
        resolved.append(path.name)
    pdfmetrics.registerFontFamily('Body', normal='Body', bold='Body-Bold',
                                 italic='Body-Italic', boldItalic='Body-BoldItalic')
    symbols = windows/'seguisym.ttf' if (windows/'seguisym.ttf').exists() else linux/'DejaVuSans.ttf'
    pdfmetrics.registerFont(TTFont('Unicode', str(symbols)))
    resolved.append(symbols.name)
    return resolved


def url_for(target, source):
    if re.match(r'^[A-Za-z][A-Za-z0-9+.-]*:', target):
        return target
    if target.startswith('#'):
        return GITHUB + source.relative_to(REPO).as_posix() + target
    return urljoin(GITHUB + quote(source.relative_to(REPO).as_posix()), target)


def clean(text):
    text = LINK_PATTERN.sub(lambda m: m.group(2), text)
    text = text.replace('**', '').replace('`', '')
    return text


def inline(text, source):
    escaped = html.escape(text)
    def hyperlink(match):
        href = html.escape(url_for(html.unescape(match.group(3)), source), quote=True)
        return f'<link href="{href}" color="#0b559f">{match.group(2)}</link>'
    escaped = LINK_PATTERN.sub(hyperlink, escaped)
    escaped = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', escaped)
    escaped = re.sub(r'`([^`]+)`', r'<font name="Code">\1</font>', escaped)
    return symbolize(escaped)


def symbolize(text):
    # Arial and Consolas lack some scientific characters; preserve their glyphs.
    return re.sub('[⁻₀ₙ∇ₖ]', lambda match: '<font name="Unicode">'+match.group()+'</font>', text)


def parse_markdown(text):
    """Small source-preserving parser for headings, tables, code, math and lists."""
    lines, i = text.splitlines(), 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('```'):
            block, i = [], i+1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                block.append(lines[i])
                i += 1
            yield 'code', '\n'.join(block)
            i += 1
            continue
        if line == '\\[':
            block, i = [], i+1
            while i < len(lines) and lines[i].strip() != '\\]':
                block.append(lines[i].strip())
                i += 1
            value = ' '.join(block)
            if '-\\nabla^2' in value:
                value = '−∇² A = μ₀ J;    B = ∇ × A'
            elif '\\int_\\Omega' in value:
                value = '∫Ω ∇Aₖ · ∇v dV = μ₀ ∫Ωbobinas Jₖ v dV'
            yield 'math', value
            i += 1
            continue
        heading = re.match(r'^(#{1,6})\s+(.+)', line)
        if heading:
            yield 'heading', (len(heading.group(1)), heading.group(2))
            i += 1
            continue
        if line.startswith('|') and line.endswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [cell.strip() for cell in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in row):
                    rows.append(row)
                i += 1
            yield 'table', rows
            continue
        if line.startswith('!['):
            image = LINK_PATTERN.fullmatch(line)
            if image:
                yield 'image', (image.group(2), image.group(3))
                i += 1
                continue
        bullet = re.match(r'^(-|\d+\.)\s+(.+)', line)
        prefix = bullet.group(1) if bullet else None
        parts = [bullet.group(2) if bullet else line]
        i += 1
        while i < len(lines) and lines[i].strip():
            next_line = lines[i].strip()
            if re.match(r'^(#|\||```|!\[|\\\[|(?:-|\d+\.)\s)', next_line):
                break
            parts.append(next_line)
            i += 1
        yield 'list' if prefix else 'paragraph', (prefix, ' '.join(parts)) if prefix else ' '.join(parts)


def add_docx_link(paragraph, label, target):
    relation = paragraph.part.relate_to(target,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), relation)
    run, props = OxmlElement('w:r'), OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0B559F'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props)
    text = OxmlElement('w:t'); text.text = clean(label); run.append(text)
    hyperlink.append(run); paragraph._p.append(hyperlink)


def add_docx_inline(paragraph, text, source):
    position = 0
    for match in LINK_PATTERN.finditer(text):
        add_docx_plain(paragraph, text[position:match.start()])
        add_docx_link(paragraph, match.group(2), url_for(match.group(3), source))
        position = match.end()
    add_docx_plain(paragraph, text[position:])


def add_docx_plain(paragraph, text):
    for piece in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not piece:
            continue
        run = paragraph.add_run(piece[2:-2] if piece.startswith('**') else
                                piece[1:-1] if piece.startswith('`') else piece)
        run.bold = piece.startswith('**')
        if piece.startswith('`'):
            run.font.name = DOCX_CODE_FONT
            run.font.size = Pt(8.7)


def table_widths(rows, total):
    count = max(map(len, rows))
    means = []
    for column in range(count):
        lengths = [len(clean(row[column])) for row in rows if column < len(row)]
        means.append(max(1.25, min(4., math_sqrt(sum(lengths)/max(len(lengths), 1)))))
    return [total*weight/sum(means) for weight in means]


def math_sqrt(value):
    return value**.5


class Manual:
    def __init__(self):
        self.fonts = font_setup()
        self.story = []
        self.doc = Document()
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(2.1), Cm(2.0)
        section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
        normal = self.doc.styles['Normal']
        normal.font.name, normal.font.size = DOCX_BODY_FONT, Pt(10)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.14
        for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
            self.doc.styles[name].font.name = DOCX_BODY_FONT
            self.doc.styles[name].font.color.rgb = RGBColor.from_string('12364A')
        header = section.header.paragraphs[0]
        self.doc.styles['Caption'].font.name = DOCX_BODY_FONT
        header.text = 'BANCADA V3  |  PROJETO E VERIFICAÇÃO DIGITAL'
        header.style = self.doc.styles['Caption']
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run('Hardware não ensaiado  •  Página ')
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
        for run in footer.runs:
            run.font.size = Pt(8)
        styles = getSampleStyleSheet()
        self.styles = {
            'body': ParagraphStyle('ManualBody', fontName='Body', fontSize=9.6, leading=13.9,
                                   spaceAfter=7, splitLongWords=1),
            'small': ParagraphStyle('ManualSmall', fontName='Body', fontSize=8, leading=11,
                                    spaceAfter=5, textColor=colors.HexColor('#4b5563'), splitLongWords=1),
            'title': ParagraphStyle('ManualTitle', fontName='Body-Bold', fontSize=28, leading=34,
                                    spaceAfter=15, textColor=colors.HexColor('#12364a')),
            'h1': ParagraphStyle('ManualH1', fontName='Body-Bold', fontSize=18, leading=23,
                                 spaceAfter=11, spaceBefore=7, keepWithNext=True,
                                 textColor=colors.HexColor('#12364a')),
            'h2': ParagraphStyle('ManualH2', fontName='Body-Bold', fontSize=12.4, leading=17,
                                 spaceAfter=8, spaceBefore=13, keepWithNext=True,
                                 textColor=colors.HexColor('#12364a')),
            'h3': ParagraphStyle('ManualH3', fontName='Body-Bold', fontSize=10.5, leading=14,
                                 spaceAfter=6, spaceBefore=8, keepWithNext=True),
            'code': ParagraphStyle('ManualCode', fontName='Code', fontSize=7.8, leading=10.3,
                                   spaceAfter=4, leftIndent=8, rightIndent=8, splitLongWords=1,
                                   backColor=colors.HexColor('#f2f5f7')),
            'table': ParagraphStyle('ManualCell', fontName='Body', fontSize=7.8, leading=10.6,
                                    splitLongWords=1),
        }
        self.width = A4[0] - 44*mm
        self.image_inserted = False

    def paragraph(self, value, source, style='body'):
        self.story.append(Paragraph(inline(value, source), self.styles[style]))
        para = self.doc.add_paragraph()
        add_docx_inline(para, value, source)
        if style == 'small':
            for run in para.runs:
                run.font.size = Pt(8)

    def heading(self, value, level=1):
        paragraph = Paragraph(html.escape(clean(value)), self.styles[f'h{min(level,3)}'])
        if level == 1:
            paragraph.manual_bookmark = clean(value)
        self.story.append(paragraph)
        self.doc.add_heading(clean(value), min(level, 3))

    def page_break(self):
        self.story.append(PageBreak())
        self.doc.add_page_break()

    def table(self, rows, source):
        ncols = max(map(len, rows))
        normalized = [row + ['']*(ncols-len(row)) for row in rows]
        data = [[Paragraph(inline(cell, source), self.styles['table']) for cell in row] for row in normalized]
        table = LongTable(data, colWidths=table_widths(rows, self.width), repeatRows=1,
                          splitByRow=1, splitInRow=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e8f1f5')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#b9c8d2')),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        self.story += [table, Spacer(1, 8)]
        word_table = self.doc.add_table(rows=0, cols=ncols)
        word_table.style = 'Table Grid'
        word_table.autofit = False
        widths = table_widths(rows, 16.6)
        for column, width in zip(word_table.columns, widths):
            column.width = Cm(width)
        for index, values in enumerate(normalized):
            cells = word_table.add_row().cells
            for column, (cell, value) in enumerate(zip(cells, values)):
                cell.width = Cm(widths[column])
                add_docx_inline(cell.paragraphs[0], value, source)
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(8)
                    if index == 0:
                        run.bold = True
                cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            if index == 0:
                repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'), 'true')
                word_table.rows[0]._tr.get_or_add_trPr().append(repeat)
        self.doc.add_paragraph()

    def code(self, value):
        for raw in value.splitlines():
            # Explicit continuation line wrapping keeps long shell commands printable.
            wrapped = textwrap.wrap(raw.expandtabs(4), width=100, replace_whitespace=False,
                                    drop_whitespace=False, break_long_words=True, break_on_hyphens=False) or [' ']
            for line in wrapped:
                self.story.append(Paragraph(symbolize(html.escape(line).replace(' ', '&#160;')), self.styles['code']))
            # Keep one Word paragraph per original command; visual wrapping must
            # not insert extra command terminators into copied DOCX text.
            para = self.doc.add_paragraph()
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.line_spacing = 1
            run = para.add_run(raw); run.font.name = DOCX_CODE_FONT; run.font.size = Pt(8)

    def image(self, caption, path):
        if path.suffix.lower() != '.png':
            return
        image = Image(str(path))
        image.drawWidth = self.width
        image.drawHeight = self.width * image.imageHeight/image.imageWidth
        self.story += [image, Paragraph(html.escape(caption), self.styles['small'])]
        self.doc.add_picture(str(path), width=Cm(16.4))
        para = self.doc.add_paragraph(caption, 'Caption')
        self.image_inserted = True

    def append_source(self, relative, title):
        source = BASE/relative
        text = source.read_text(encoding='utf-8-sig')
        self.page_break(); self.heading(title)
        self.paragraph(f'Fonte consolidada: [{relative}]({relative}). Conteúdo técnico integral; documentos Markdown permanecem a referência editável.', BASE/'README.md', 'small')
        for kind, value in parse_markdown(text):
            if kind == 'heading':
                level, heading = value
                if level != 1:
                    self.heading(heading, min(level, 3))
                else:
                    self.paragraph(heading, source, 'small')
            elif kind == 'paragraph':
                self.paragraph(value, source)
            elif kind == 'list':
                prefix, content = value
                self.paragraph(('• ' if prefix == '-' else prefix+' ') + content, source)
            elif kind == 'table':
                self.table(value, source)
            elif kind == 'code':
                self.code(value)
            elif kind == 'math':
                self.paragraph(value, source)
            elif kind == 'image':
                caption, target = value
                self.image(caption, (source.parent/target).resolve())

    def save(self):
        pdf = BASE/'Manual_Bancada_V3.pdf'
        docx = BASE/'Manual_Bancada_V3.docx'
        document = ManualDocTemplate(str(pdf), pagesize=A4, rightMargin=22*mm,
            leftMargin=22*mm, topMargin=24*mm, bottomMargin=21*mm,
            title='Manual Bancada V3 — projeto e verificação digital',
            author='Projeto Núcleo de Campo Toroidal — 7dsol3',
            subject='Hardware não ensaiado; eletrônica, aquisição, controle e FEM preliminar')
        def furniture(canvas, doc):
            canvas.setFont('Body', 7.7); canvas.setFillColor(colors.HexColor('#52616b'))
            canvas.drawString(22*mm, A4[1]-14*mm, 'BANCADA V3  |  PROJETO E VERIFICAÇÃO DIGITAL')
            canvas.drawString(22*mm, 12*mm, 'Hardware não ensaiado — resultados calculados não são medições')
            canvas.drawRightString(A4[0]-22*mm, 12*mm, f'Página {doc.page}')
        document.build(self.story, onFirstPage=furniture, onLaterPages=furniture)
        self.doc.core_properties.title = 'Manual Bancada V3 — projeto e verificação digital'
        self.doc.core_properties.subject = 'Hardware não ensaiado'
        self.doc.core_properties.author = 'Projeto Núcleo de Campo Toroidal — 7dsol3'
        self.doc.save(docx)
        return pdf, docx


def validate(pdf, docx, source_paths):
    with fitz.open(pdf) as document:
        all_text = '\n'.join(page.get_text() for page in document)
        required = ['HARDWARE NÃO ENSAIADO', 'CA principal', 'CC de mapa/FEM',
                    'PI CC', 'Toroide auxiliar', '57,7779', '2,79%', '11,07%',
                    'OPA548', 'DRV425', 'OPA4277', 'INCONCLUSIVE', 'MCP4725',
                    '1 MSps', '4,70', '500 kHz', 'Ensaio CC estacionário',
                    'analisar_cc.py', 'metadados_cc_pendente.json']
        for phrase in required:
            if phrase not in all_text:
                raise AssertionError(f'PDF required text missing: {phrase}')
        outside = []
        blank_pages = []
        for number, page in enumerate(document, 1):
            if len(page.get_text().strip()) < 100:
                blank_pages.append(number)
            for word in page.get_text('words'):
                x0, y0, x1, y1 = word[:4]
                if x0 < 0 or y0 < 0 or x1 > page.rect.width+.5 or y1 > page.rect.height+.5:
                    outside.append((number, word[4]))
        assert not outside, ('Text outside PDF page', outside[:10])
        assert not blank_pages, ('Blank PDF pages', blank_pages)
        assert '\ufffd' not in all_text and '\x00' not in all_text
        links = sum(len(page.get_links()) for page in document)
        assert links >= 80, f'Expected original clickable references, got {links}'
        bookmarks = len(document.get_toc())
        assert bookmarks >= 14, f'Expected chapter navigation, got {bookmarks}'
        pages = len(document)
        characters = len(all_text)
    with zipfile.ZipFile(docx) as container:
        assert container.testzip() is None
        assert '[Content_Types].xml' in container.namelist()
        assert 'word/document.xml' in container.namelist()
        assert any(name.startswith('word/media/') for name in container.namelist())
        xml = container.read('word/document.xml').decode('utf-8')
        assert '<w:document' in xml and 'HARDWARE NÃO ENSAIADO' in xml
        assert '<!DOCTYPE' not in xml
    reopened = Document(docx)
    assert len(reopened.tables) >= 15
    return {'pdf_pages': pages, 'pdf_text_characters': characters,
            'pdf_clickable_links': links, 'pdf_text_outside_page': 0,
            'pdf_chapter_bookmarks': bookmarks,
            'docx_tables': len(reopened.tables), 'docx_paragraphs': len(reopened.paragraphs),
            'source_documents': len(source_paths),
            'source_sha256': {str(p.relative_to(BASE)): hashlib.sha256(p.read_bytes()).hexdigest()
                              for p in source_paths},
            'hardware_tested': False, 'measured_results_created': False}


def main():
    source_paths = [BASE/path for path, _ in SOURCES]
    missing = [str(path) for path in source_paths if not path.is_file()]
    if missing:
        raise RuntimeError('Wait for complete V3 sources before building: '+', '.join(missing))
    manual = Manual()
    manual.story.append(Spacer(1, 28*mm))
    manual.story.append(Paragraph('MANUAL DA<br/>BANCADA V3', manual.styles['title']))
    manual.doc.add_heading('MANUAL DA BANCADA V3', 0)
    manual.heading('PROJETO E VERIFICAÇÃO DIGITAL', 2)
    manual.heading('HARDWARE NÃO ENSAIADO', 2)
    manual.paragraph('Núcleo eletromagnético de bancada: projeto elétrico, aquisição sincronizada, calibração, controle CC e FEM 3D preliminar.', BASE/'README.md')
    manual.paragraph('Esta edição consolida os documentos técnicos completos do projeto. Nenhuma montagem, calibração, captura experimental ou qualificação térmica/elétrica é apresentada como realizada. Software e resultados calculados não substituem ensaio físico.', BASE/'README.md')
    manual.paragraph('Referência digital: [repositório 7dsol3/nucleo-campo-toroidal-quantico](https://github.com/7dsol3/nucleo-campo-toroidal-quantico).', BASE/'README.md')
    manual.paragraph('O erro de 2,79% refere-se ao ponto central do caso FEM e às hipóteses documentadas; não é limite de erro certificado para todo o domínio 3D.', BASE/'README.md')
    manual.page_break(); manual.heading('Roteiro de consulta e estado da evidência')
    manual.table([
        ['Modo', 'Configuração e comparação permitida'],
        ['CA principal', 'Duas fontes sincronizadas, capacitores em série, receptor com carga; 2 V RMS, 1.400 Hz e fase 0° como primeiro ponto a qualificar.'],
        ['CC de mapa/FEM (CC2)', 'Duas bobinas, 0,150 A CC medidos em cada uma, capacitores contornados e receptor aberto; corresponde ao FEM estático publicado.'],
        ['PI CC (PI1)', 'Uma bobina controlada pelo Pico/MCP4725; segunda fonte e receptor desconectados; não reproduz automaticamente o mapa FEM de duas bobinas.'],
        ['Toroide auxiliar', 'Ensaio separado com alimentação e potência próprias; sem terceiro driver CA simultâneo implementado.'],
    ], BASE/'README.md')
    for _, title in SOURCES:
        manual.paragraph(title, BASE/'README.md')
    manual.paragraph('Anexos: lista de materiais elétricos integral e componentes do filtro. Os arquivos mecânicos CAD/STL continuam na V2. Os links de arquivos abrem sua localização no GitHub; links de fabricantes preservam os destinos originais.', BASE/'README.md')
    for path, title in SOURCES:
        manual.append_source(path, title)
    for rel, title in [('eletronica/BOM.csv', 'Anexo A — lista integral de materiais elétricos'),
                       ('aquisicao/filtro_componentes.csv', 'Anexo B — componentes do filtro analógico')]:
        path = BASE/rel
        source_paths.append(path)
        with path.open(encoding='utf-8-sig', newline='') as stream:
            records = list(csv.DictReader(stream))
        manual.page_break(); manual.heading(title)
        manual.paragraph('Transcrição integral das linhas do CSV. Quantidades, referências e estado de compra continuam condicionados às verificações descritas no projeto.', path)
        for record in records:
            identifier = record.get('ref') or record.get('reference') or next(iter(record.values()))
            manual.heading(str(identifier), 3)
            rows = [['Campo do CSV', 'Valor original']]
            for key, value in record.items():
                if value and str(value).startswith('http'):
                    value = f'[Referência oficial]({value})'
                rows.append([key, value or ''])
            manual.table(rows, path)
    assert manual.image_inserted, 'Expected calculated FEM PNG in manual'
    pdf, docx = manual.save()
    result = validate(pdf, docx, source_paths)
    result['embedded_pdf_fonts'] = manual.fonts
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
