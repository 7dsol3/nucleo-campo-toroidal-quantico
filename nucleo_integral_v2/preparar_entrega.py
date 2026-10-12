"""Gera documentos e BOM da versão integral. Preserva CAD e cálculos."""
from pathlib import Path
import csv,json,html,re
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,Preformatted,PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Inches,Pt

ROOT=Path(__file__).resolve().parent

def parse(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('```'):
            i+=1;data=[]
            while i<len(lines) and not lines[i].strip().startswith('```'):data.append(lines[i]);i+=1
            yield'code','\n'.join(data);i+=1;continue
        if line.startswith('|'):
            data=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch('[-: ]+',v) for v in row):data.append(row)
                i+=1
            yield'table',data;continue
        if line.startswith('#'):yield'h'+str(len(line)-len(line.lstrip('#'))),line.lstrip('# ').strip();i+=1;continue
        yield'p',line;i+=1

def prepare():
    modules=json.loads((ROOT/'pesquisa/modulos.json').read_text(encoding='utf-8'))
    (ROOT/'pesquisa/modulos.js').write_text('window.NUCLEO_MODULES='+json.dumps(modules,ensure_ascii=False,separators=(',',':'))+';',encoding='utf-8')
    scene=json.loads((ROOT/'cad/scene.json').read_text(encoding='utf-8'))
    p=ROOT/'laboratorio.js';s=p.read_text(encoding='utf-8')
    s=re.sub(r'\d+ elementos de referência',str(len(scene['parts']))+' elementos de referência',s)
    p.write_text(s,encoding='utf-8')
    p=ROOT/'index.html';s=p.read_text(encoding='utf-8').replace('Soma ativa assinada das fontes AC','Duas portas axiais • potência ativa');p.write_text(s,encoding='utf-8')
    p=ROOT/'circuito_integral.svg';s=p.read_text(encoding='utf-8').replace('10 Ω / 5 W','10 Ω / 10 W');p.write_text(s,encoding='utf-8')
    for name in ['Projeto_Integral.md','LEIA_PRIMEIRO.md']:
        p=ROOT/name;s=p.read_text(encoding='utf-8')
        s=re.sub(r'(?<=\d)(?=(?:mm|mV|mA|mH|µH|µT|µW|V|Ω|A|W|Hz|ms|°C)\b)',' ',s)
        p.write_text(s,encoding='utf-8')
    with (ROOT/'cad/pecas.csv').open(encoding='utf-8-sig') as f:mechanical=list(csv.DictReader(f,delimiter=';'))
    rows=[{'codigo':r['codigo'],'item':r['peca'],'quantidade':r['quantidade'],'especificacao':r['material'],'observacao':r['fabricacao'],'arquivo':r['arquivo']} for r in mechanical]
    parts=[
        ('D01','Capacitor filme não polarizado','2','10 µF ±5%; nominal≥100 VDC','Conferir limite AC, corrente RMS/frequência e ESR; não é eletrolítico comum'),
        ('D02','Resistor não indutivo','2','10 Ω / 10 W','Medir resistência/deriva e resposta AC; usado para referência de corrente'),
        ('D03','Carga receptora','1','100 Ω / ≥0,25 W, não indutivo','Medir valor real e tensão/corrente; modelo usa0,25 W como limite'),
        ('D04','Fonte axial de dois canais sincronizados','1 conjunto','100–5000 Hz;0–5 V RMS por porta;limite0,5 A RMS','Saídas e sondas compatíveis com isolamento/topologia; fase0–180°; potência de retorno conforme equipamento'),
        ('D05','Fonte toroidal independente','1','Senoidal limitada;0–0,5 A RMS; faixa100–5000 Hz','Corrente controlada no toroide de baixa indutância; medir tensão/perdas/fuga'),
        ('D06','Fonte DC para testes','1 conjunto','Isolada;saídas iguais,limitadas;0–5 V','MapaDC e PI com capacitores contornados,receptor desconectado'),
        ('D07','Cabos/conectores e proteção','1 conjunto','Compatíveis com correntes,tensões e topologia do ensaio','Terminais isolados;parada;instrumentar os limites'),
        ('D08','Instrumentação V/I/B/T','1 conjunto','Conforme pesquisa/requisitos_instrumentos.csv','Sondas diferenciais/isoladas, campo calibradoDC/AC,LCR,termômetro,DAQ ou osciloscópio'),
        ('D09','Placa de controle PI','1 proposta','ADC/DAC e fontes/driver com desempenho caracterizado','Não há PCB/firmware/hardware validado: gráfico PI é ensaio numérico separado'),
    ]
    rows += [dict(codigo=c,item=n,quantidade=q,especificacao=spec,observacao=note,arquivo='') for c,n,q,spec,note in parts]
    with (ROOT/'lista_completa_pecas.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=['codigo','item','quantidade','especificacao','observacao','arquivo'],delimiter=';');w.writeheader();w.writerows(rows)
    return rows

def documents(rows):
    pdfmetrics.registerFont(TTFont('IntegralArial',r'C:\Windows\Fonts\arial.ttf'))
    pdfmetrics.registerFont(TTFont('IntegralBold',r'C:\Windows\Fonts\arialbd.ttf'))
    pdfmetrics.registerFont(TTFont('IntegralMono',r'C:\Windows\Fonts\consola.ttf'))
    styles=getSampleStyleSheet()
    for name,font,size,leading,color in [('Body','IntegralArial',9.1,13,'#283c49'),('Head','IntegralBold',14,19,'#174f61'),('Title','IntegralBold',25,30,'#174f61'),('Cell','IntegralArial',7.7,10.5,'#263846'),('Code','IntegralMono',7.6,10,'#263846')]:
        styles.add(ParagraphStyle(name='I'+name,fontName=font,fontSize=size,leading=leading,textColor=colors.HexColor(color),spaceAfter=7 if name in('Body','Title') else 0,spaceBefore=14 if name=='Head' else 0,keepWithNext=name=='Head'))
    def para(text,style='IBody'):return Paragraph(html.escape(text),styles[style])
    def table(data,widths):
        t=Table([[para(v,'ICell') for v in row] for row in data],colWidths=widths,repeatRows=1)
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#d8e8ef')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f2f5f7')]),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#c1ced5')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));return t
    flow=[para('Núcleo integral','ITitle'),para('Projeto de engenharia, simulação 3D e matriz de ensaios — V2'),para('Referência Adilson Oliveira / 7dsolv • 12/10/2026 • protótipo físico ainda não construído')]
    pic=ROOT/'cad/previa_3d.png'
    if pic.exists():flow.extend([Image(str(pic),width=470,height=376),para('Montagem CAD desta entrega. Dimensões propostas: base360 mm; aros340 mm; bobinas/receptor/toroide e instrumentação em ar comum.')])
    doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
    doc.sections[0].header.paragraphs[0].text='7DSOLV / NÚCLEO INTEGRAL / V2'
    doc.sections[0].footer.paragraphs[0].text='Projeto e cálculos • 12/10/2026 • sem ensaio físico'
    doc.add_heading('Núcleo integral',0);doc.add_paragraph('Versão2 —12/10/2026. Referência visual: Adilson Oliveira /7dsolv. Cotas e circuitos são propostas desta entrega.')
    if pic.exists():doc.add_picture(str(pic),width=Inches(5.7))
    for kind,data in parse((ROOT/'Projeto_Integral.md').read_text(encoding='utf-8')):
        if kind=='h1':continue
        if kind.startswith('h'):
            flow.append(para(data,'IHead'));doc.add_heading(data,level=min(int(kind[1:])-1,3))
            if data.startswith('5. Modelo') and (ROOT/'Circuito_Integral.png').exists():
                flow.extend([Image(str(ROOT/'Circuito_Integral.png'),width=515,height=305),Spacer(1,12)]);doc.add_picture(str(ROOT/'Circuito_Integral.png'),width=Inches(6.1))
            continue
        if kind=='table':
            widths=[210,305] if len(data[0])==2 else [145,185,185]
            flow.extend([table(data,widths),Spacer(1,10)])
            dt=doc.add_table(rows=0,cols=len(data[0]));dt.style='Light Shading Accent 1'
            for row in data:
                cells=dt.add_row().cells
                for i,v in enumerate(row):cells[i].text=v
            continue
        if kind=='code':
            flow.append(Preformatted(data,styles['ICode'],maxLineLength=96));dp=doc.add_paragraph();r=dp.add_run(data);r.font.name='Consolas';r.font.size=Pt(8);continue
        flow.append(para(data));doc.add_paragraph(data)
    flow.extend([PageBreak(),para('Apêndice — lista integral de peças','IHead'),para('A especificação completa com materiais, fabricação e arquivos está em lista_completa_pecas.csv e cad/LEIA-ME_CAD.md. Quantidades são de projeto; nenhuma compra ou fabricação foi realizada.')])
    data=[['Código','Peça','Qtd.','Material/requisito']]+[[r['codigo'],r['item'],r['quantidade'],r['especificacao']] for r in rows]
    flow.append(table(data,[40,230,50,195]));doc.add_page_break();doc.add_heading('Lista integral de peças',1)
    dt=doc.add_table(rows=0,cols=4);dt.style='Light Shading Accent 1'
    for row in data:
        cells=dt.add_row().cells
        for i,v in enumerate(row):cells[i].text=v
    def footer(c,d):
        c.saveState();c.setFont('IntegralArial',7.5);c.setFillColor(colors.HexColor('#536c79'));c.drawString(40,28,'7DSOLV / V2 • previsões numéricas • sem ensaio físico');c.drawRightString(555,28,'Página '+str(d.page));c.restoreState()
    SimpleDocTemplate(str(ROOT/'Projeto_Integral.pdf'),pagesize=A4,leftMargin=40,rightMargin=40,topMargin=40,bottomMargin=50,title='Núcleo integral — Projeto de engenharia e ensaios',author='Referência Adilson Oliveira / 7dsolv').build(flow,onFirstPage=footer,onLaterPages=footer)
    doc.save(ROOT/'Projeto_Integral.docx')
    print('Documentos integrais e BOM gerados.')

if __name__=='__main__':documents(prepare())
