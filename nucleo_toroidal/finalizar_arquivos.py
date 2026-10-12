"""Gera documentos e BOM da entrega local; não altera o modelo físico."""
from pathlib import Path
import csv
import html
import re
import json
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, Preformatted, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt

ROOT=Path(__file__).resolve().parent

def adjust_viewer():
    p=ROOT/'index.html'
    s=p.read_text(encoding='utf-8')
    replacements=[
        ('≈ 65 m','≈ 70 m'),('height/285','height/330'),('height*.55-Z','height*.53-Z'),
        ('z-7.7,z-5.7,95.25,110','z-7.7,z-5.7,94,110'),
        ('z+5.7,z+7.7,95.25,110','z+5.7,z+7.7,94,110'),
        ('annulus(x,y,67.7,60+state.separation-7.7','annulus(x,y,low+7.7,low+7.7+state.separation-15.4'),
        ('annulus(x,y,-4,176,0,1.5','annulus(x,y,-5,Math.max(175,60+state.separation+15),0,1.5'),
        ('pctx.fillText(fmt(max,0),barx-3,top-9)',"pctx.fillText('≥'+fmt(max,0),barx-8,top-9)"),
        ('Biot–Savart numérico • cores: módulo RMS • setas: sentido no pico positivo','Biot–Savart • RMS • escala automática recortada no percentil 98; ≥ indica saturação de cor'),
        ('Circuito RLC série • curva prevista com os parâmetros atuais','Circuito RLC série • corrente normalizada por tensão terminal, I/V'),
        ('imax=state.voltage/p.R','imax=1/p.R'),
        ("pctx.fillText('I RMS / mA',left,top-22)","pctx.fillText('I / V • mA por volt RMS',left,top-22)"),
        ('iy(circuit(f).I)','iy(circuit(f,1).I)'),('iy(p.I)','iy(circuit(state.frequency,1).I)'),
        ('if(m.frequency_hz<100||m.frequency_hz>5000)continue','if(m.frequency_hz<100||m.frequency_hz>5000||m.source_rms_v<=0)continue'),
        ('iy(m.current_rms_a)','iy(m.current_rms_a/m.source_rms_v)'),
        ('● Medições importadas (V informado no CSV)','● Medições I/V (linhas de 0 V omitidas)'),
    ]
    for before,after in replacements:
        s=s.replace(before,after)
    p.write_text(s,encoding='utf-8')

def bom():
    mechanical=list(csv.DictReader((ROOT/'cad/pecas.csv').open(encoding='utf-8-sig'),delimiter=';'))
    rows=[{'codigo':r['codigo'],'item':r['peca'],'quantidade':r['quantidade'],'especificacao':r['medidas_mm'],'observacao':r['fabricacao_observacao']} for r in mechanical]
    additions=[
        ('E03','Capacitor de filme','1','10 µF ±5%; polipropileno não polarizado; ≥100 VDC','Verificar tensão AC e corrente/frequência no catálogo; medir C/ESR'),
        ('E04','Resistor não indutivo','1','10 Ω / 10 W','Medir R; observar condição de dissipação do fabricante'),
        ('E05','Fonte senoidal isolada','1','100–5000 Hz; saída0–5 V RMS; limite≤0,5 A RMS','Equipamento externo ainda a selecionar; não há PCB de amplificador nesta entrega'),
        ('E06','Fonte DC de bancada','1','Isolada, baixa tensão, limite até0,5 A','Para mapa estático; capacitor série removido'),
        ('E07','Cabos e terminais isolados','1 conjunto','Compatíveis com até0,5 A RMS; conexão de baixa resistência','Modelo reserva1m de terminais além das espiras; medir Rreal'),
        ('I01','Osciloscópio e sondas','1 conjunto','Medida simultânea de V/I e fase; banda adequada','Usar sondas diferenciais/isoladas conforme a saída da fonte'),
        ('I02','Medidor LCR / multímetro','1 conjunto','L/C/R em frequência; true RMS adequado até5kHz','Conferir calibração e especificação de banda'),
        ('I03','Sensor de campo DC/AC','1 conjunto','Resolução desejada~1µT; calibrado até5kHz para AC','Pode usar HallDC e bobina de buscaAC separadamente'),
        ('I04','Termômetro','1','Medição do enrolamento e suportes','Registrar T; parada inicial45°C até validar materiais'),
        ('I05','Gabarito e instrumentos dimensionais','1 conjunto','Suporte não magnético, paquímetro/régua','Posicionamento desejado~1mm'),
        ('F01','Fabricação','1 serviço','Mesa3D≥230mm útil ou serviço de impressão; corte/furação da base','Carretel exige suporte removível; conferir encaixe físico'),
    ]
    rows += [dict(zip(['codigo','item','quantidade','especificacao','observacao'],x)) for x in additions]
    with (ROOT/'lista_completa_pecas.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=['codigo','item','quantidade','especificacao','observacao'],delimiter=';');writer.writeheader();writer.writerows(rows)

def parse_md(text):
    lines=text.splitlines(); i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('```'):
            i+=1; block=[]
            while i<len(lines) and not lines[i].strip().startswith('```'):
                block.append(lines[i]);i+=1
            yield ('code','\n'.join(block));i+=1;continue
        if line.startswith('|'):
            data=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',x) for x in row):data.append(row)
                i+=1
            yield('table',data);continue
        if line.startswith('#'):
            yield ('h'+str(len(line)-len(line.lstrip('#'))),line.lstrip('# ').strip());i+=1;continue
        yield('p',line);i+=1

def generate_docs():
    pdfmetrics.registerFont(TTFont('LabArial',r'C:\Windows\Fonts\arial.ttf'))
    pdfmetrics.registerFont(TTFont('LabArialBold',r'C:\Windows\Fonts\arialbd.ttf'))
    pdfmetrics.registerFont(TTFont('LabMono',r'C:\Windows\Fonts\consola.ttf'))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='LabBody',fontName='LabArial',fontSize=9.2,leading=13,spaceAfter=7,textColor=colors.HexColor('#202e3a')))
    styles.add(ParagraphStyle(name='LabHead',fontName='LabArialBold',fontSize=14,leading=18,spaceBefore=16,spaceAfter=8,keepWithNext=True,textColor=colors.HexColor('#164b59')))
    styles.add(ParagraphStyle(name='LabCell',fontName='LabArial',fontSize=8,leading=11,spaceAfter=0))
    styles.add(ParagraphStyle(name='LabCode',fontName='LabMono',fontSize=7.5,leading=10,spaceAfter=10))
    styles.add(ParagraphStyle(name='LabTitle',fontName='LabArialBold',fontSize=24,leading=29,spaceAfter=14,textColor=colors.HexColor('#164b59')))
    def para(text,style='LabBody'):
        esc=html.escape(text)
        esc=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',esc)
        esc=re.sub(r'`(.*?)`',r'<font name="LabMono">\1</font>',esc)
        return Paragraph(esc,styles[style])
    flow=[para('Núcleo — bancada e simulação 3D','LabTitle'),para('Projeto preliminar de engenharia • referência Adilson Oliveira / 7dsolv • 12/10/2026'),para('Campo magnético, ressonância elétrica e medidas de bancada. Cálculos e arquivos digitais; protótipo físico ainda não construído.')]
    pic=ROOT/'cad/previa_3d.png'
    if pic.exists():
        flow += [Image(str(pic),width=430,height=344),para('Renderização da montagem CAD real desta entrega. Dois anéis de bobinas circulares, suportes isolantes e base; sem feixes ou efeitos quânticos atribuídos.')]
    doc=Document(); section=doc.sections[0];section.header.paragraphs[0].text='7DSOLV / LABORATÓRIO DE BOBINAS'
    section.footer.paragraphs[0].text='Projeto preliminar • 12/10/2026 • sem validação experimental nesta entrega'
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    doc.add_heading('Núcleo — bancada e simulação 3D',0)
    doc.add_paragraph('Referência autoral: Adilson Oliveira / 7dsolv. Desenvolvimento preliminar: 12/10/2026.')
    doc.add_paragraph('Campo magnético, ressonância e balanço de potência; protótipo físico ainda não construído.')
    if pic.exists():doc.add_picture(str(pic),width=Inches(5.2))
    for kind,data in parse_md((ROOT/'Projeto_Tecnico.md').read_text(encoding='utf-8')):
        if kind=='h1':continue
        if kind.startswith('h'):
            flow.append(para(data,'LabHead'));doc.add_heading(data,level=min(int(kind[1:])-1,3))
            if data.startswith('4. Circuito') and (ROOT/'Circuito_Ensaio.png').exists():
                flow.extend([Image(str(ROOT/'Circuito_Ensaio.png'),width=515,height=220),Spacer(1,10)])
                doc.add_picture(str(ROOT/'Circuito_Ensaio.png'),width=Inches(6.1))
            continue
        if kind=='p':
            flow.append(para(data));doc.add_paragraph(re.sub(r'[*`]', '',data));continue
        if kind=='code':
            flow.append(Preformatted(data,styles['LabCode'],maxLineLength=96))
            pp=doc.add_paragraph();run=pp.add_run(data);run.font.name='Consolas';run.font.size=Pt(8);continue
        if kind=='table':
            ncols=len(data[0]);widths=([115,70,330] if ncols==3 else [210,305])
            if ncols!=len(widths):widths=[515/ncols]*ncols
            cells=[[para(cell,'LabCell') for cell in row] for row in data]
            table=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dcebef')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f3f6f7')]),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#c3d2d8')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
            flow.extend([table,Spacer(1,10)])
            t=doc.add_table(rows=0,cols=ncols);t.style='Light Shading Accent 1'
            for row in data:
                cells=t.add_row().cells
                for j,val in enumerate(row):cells[j].text=val
    def footer(c,d):
        c.saveState();c.setStrokeColor(colors.HexColor('#c3d2d8'));c.line(40,40,555,40);c.setFont('LabArial',7.5);c.setFillColor(colors.HexColor('#536b76'));c.drawString(40,27,'7DSOLV • Projeto preliminar • sem ensaio físico');c.drawRightString(555,27,'Página '+str(d.page));c.restoreState()
    pdf=SimpleDocTemplate(str(ROOT/'Projeto_Tecnico.pdf'),pagesize=A4,rightMargin=40,leftMargin=40,topMargin=38,bottomMargin=52,title='Núcleo — Projeto de bancada e simulação 3D',author='Referência Adilson Oliveira / 7dsolv; engenharia preliminar')
    pdf.build(flow,onFirstPage=footer,onLaterPages=footer)
    doc.save(ROOT/'Projeto_Tecnico.docx')

if __name__=='__main__':
    adjust_viewer();bom();generate_docs()
    print('Documentos e lista completa de peças gerados.')
